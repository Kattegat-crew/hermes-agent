#!/usr/bin/env python3
"""Manual de Asistentes Inteligentes — formato oficial NeuralCrew Labs (membrete aprobado).
Recrea el contenido del PDF adjunto con la identidad de marca oficial.
"""
import os, io
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

GOLD = RGBColor(0xD4, 0xAF, 0x37)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
DGRAY = RGBColor(0x3A, 0x3A, 0x3A)
MGRAY = RGBColor(0x70, 0x70, 0x70)
FONT = "Calibri"

ICON = "/root/hermes-agent/data/workspace/logo_icon.png"
OUT = "/root/hermes-agent/data/workspace/Manual_Asistentes_Inteligentes_NeuralCrew_Labs.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
sec.top_margin = Cm(1.7)
sec.bottom_margin = Cm(1.6)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)

normal = doc.styles['Normal']
normal.font.name = FONT
normal.font.size = Pt(10.5)
normal.font.color.rgb = BLACK

def gold_bottom_border(paragraph, sz=8):
    paragraph._p.get_or_add_pPr().append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="2" w:color="D4AF37"/></w:pBdr>'))

def clean(tbl):
    tblPr = tbl._tbl.tblPr
    tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none" w:sz="0"/><w:left w:val="none" w:sz="0"/>'
        f'<w:bottom w:val="none" w:sz="0"/><w:right w:val="none" w:sz="0"/>'
        f'<w:insideH w:val="none" w:sz="0"/><w:insideV w:val="none" w:sz="0"/>'
        f'</w:tblBorders>'))

def cb(cell):
    cell._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none" w:sz="0"/><w:left w:val="none" w:sz="0"/>'
        f'<w:bottom w:val="none" w:sz="0"/><w:right w:val="none" w:sz="0"/>'
        f'</w:tcBorders>'))

def t(rows, cols, widths=None, align_center=False):
    tbl = doc.add_table(rows=rows, cols=cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER if align_center else WD_TABLE_ALIGNMENT.LEFT
    try:
        tbl.autofit = False
    except Exception:
        pass
    tblPr = tbl._tbl.tblPr
    tblPr.append(parse_xml(
        f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    tblPr.append(parse_xml(
        f'<w:tblW {nsdecls("w")} w:w="{int(sum(widths)*567) if widths else 9900}" w:type="dxa"/>'))
    tblPr.append(parse_xml(
        f'<w:tblCellMar {nsdecls("w")}>'
        f'<w:top w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/>'
        f'<w:bottom w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/>'
        f'</w:tblCellMar>'))
    clean(tbl)
    for r in tbl.rows:
        for i, c in enumerate(r.cells):
            cb(c)
            if widths and i < len(widths):
                c.width = Cm(widths[i])
                tcPr = c._tc.get_or_add_tcPr()
                tcPr.append(parse_xml(
                    f'<w:tcW {nsdecls("w")} w:w="{int(widths[i]*567)}" w:type="dxa"/>'))
    return tbl

# ══ LETTERHEAD ═══════════════════════════════════════
# top gold rule (thin, near top)
ptop = doc.add_paragraph()
ptop.paragraph_format.space_after = Pt(2)
ptop._p.get_or_add_pPr().append(parse_xml(
    f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="12" w:space="1" w:color="D4AF37"/></w:pBdr>'))

hdr = t(1, 2, [2.8, 13.4])
icon_cell, name_cell = hdr.rows[0].cells

p_icon = icon_cell.paragraphs[0]
p_icon.alignment = WD_ALIGN_PARAGRAPH.CENTER
if os.path.exists(ICON):
    with Image.open(ICON) as im:
        im = im.convert("RGBA"); im.thumbnail((400,400), Image.Resampling.LANCZOS)
        buf = io.BytesIO(); im.save(buf,"PNG"); buf.seek(0)
        p_icon.add_run().add_picture(buf, height=Cm(2.1))

p_n1 = name_cell.paragraphs[0]
p_n1.paragraph_format.space_before = Pt(3); p_n1.paragraph_format.space_after = Pt(0)
r = p_n1.add_run("NEURALCREW"); r.bold=True; r.font.size=Pt(22); r.font.color.rgb=BLACK; r.font.name=FONT

p_n2 = name_cell.add_paragraph()
p_n2.paragraph_format.space_before = Pt(0); p_n2.paragraph_format.space_after = Pt(2)
r = p_n2.add_run("LABS"); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=GOLD; r.font.name=FONT
gold_bottom_border(p_n2, 4)

p_tag = name_cell.add_paragraph()
p_tag.paragraph_format.space_before = Pt(3); p_tag.paragraph_format.space_after = Pt(0)
r = p_tag.add_run("Inteligencia artificial aplicada a tu negocio")
r.italic=True; r.font.size=Pt(9); r.font.color.rgb=DGRAY; r.font.name=FONT

# metadata row
def put(cell_label, cell_val, label, val):
    p0 = cell_label.paragraphs[0]; p0.paragraph_format.space_before=Pt(6); p0.paragraph_format.space_after=Pt(0)
    r0 = p0.add_run(label+":"); r0.bold=True; r0.font.size=Pt(7.5); r0.font.color.rgb=MGRAY
    p1 = cell_val.paragraphs[0]; p1.paragraph_format.space_before=Pt(6); p1.paragraph_format.space_after=Pt(0)
    r1 = p1.add_run(val); r1.font.size=Pt(7.5); r1.font.color.rgb=BLACK

m1 = t(1, 4, [2.4, 5.6, 2.4, 3.4])
c0,c1,c2,c3 = m1.rows[0].cells
put(c0, c1, "CLIENTE", "Golden Game / Lucky Paradise")
put(c2, c3, "FECHA", "24/08/2026")
m2 = t(1, 4, [2.4, 8.6, 2.4, 2.4])
c0,c1,c2,c3 = m2.rows[0].cells
put(c0, c1, "ASUNTO", "Manual de asistentes inteligentes")
put(c2, c3, "DOC. Nº", "NLC-0001")

# Title
tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.paragraph_format.space_before = Pt(14); tp.paragraph_format.space_after = Pt(2)
r = tp.add_run("ASISTENTE INTELIGENTE")
r.bold=True; r.font.size=Pt(19); r.font.color.rgb=BLACK; r.font.name=FONT

tp2 = doc.add_paragraph(); tp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp2.paragraph_format.space_before = Pt(0); tp2.paragraph_format.space_after = Pt(2)
r = tp2.add_run("NEURALCREW — Asistentes de WhatsApp · Equipo Directivo")
r.bold=True; r.font.size=Pt(12); r.font.color.rgb=GOLD; r.font.name=FONT
gold_bottom_border(tp2, 8)

sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_before = Pt(4); sub.paragraph_format.space_after = Pt(2)
r = sub.add_run("Guía práctica para el equipo directivo de Golden Game y Lucky Paradise")
r.italic=True; r.font.size=Pt(10.5); r.font.color.rgb=DGRAY; r.font.name=FONT

sub2 = doc.add_paragraph(); sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub2.paragraph_format.space_before = Pt(0); sub2.paragraph_format.space_after = Pt(6)
r = sub2.add_run("Helmer · Jaqueline · Yulieth")
r.font.size=Pt(10); r.font.color.rgb=MGRAY; r.font.name=FONT

# ══ BODY ══════════════════════════════════════════════
def heading(txt):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(13); h.paragraph_format.space_after = Pt(4)
    r1 = h.add_run(txt); r1.bold=True; r1.font.size=Pt(12.5); r1.font.color.rgb=BLACK; r1.font.name=FONT

def body(txt):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.15
    r = p.add_run(txt); r.font.size=Pt(10.5); r.font.color.rgb=DGRAY; r.font.name=FONT

def bullet(txt, indent=0.4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.left_indent = Cm(indent)
    rb = p.add_run("•  "); rb.font.color.rgb=GOLD; rb.font.size=Pt(10); rb.bold=True
    rt = p.add_run(txt); rt.font.size=Pt(10.5); rt.font.color.rgb=DGRAY; rt.font.name=FONT

heading("1. Qué son estos asistentes")
body("Son asistentes virtuales con inteligencia artificial que trabajan 24/7. Atienden en WhatsApp hablando en español colombiano, con el tono que corresponde a cada casino. Están conectados a la gestión del día a día: clientes (CRM), correo, calendario, documentos y reportes.")
body("Este manual está enfocado al uso por el equipo directivo. Cada persona tiene su propio asistente.")

heading("2. Qué puede hacer el asistente")
for b in [
    "Atender consultas de clientes por WhatsApp al instante.",
    "Registrar y gestionar leads en el CRM (embudo de ventas).",
    "Agendar citas y eventos en el calendario.",
    "Enviarte reportes diarios, recordatorios de citas y resúmenes semanales.",
    "Leer y redactar correos desde la bandeja de la empresa.",
    "Buscar documentos e información en Google Drive.",
    "Investigar competencia y tendencias en internet.",
    "Escuchar notas de voz y responder oralmente si lo necesitas.",
    "Analizar imágenes, capturas de pantalla y documentos escaneados.",
]:
    bullet(b)

heading("3. Cómo se usa (organización básica)")
body("No necesitas instalar nada. Funciona igual que un contacto de WhatsApp normal:")
for b in [
    "Escríbele cualquier consulta o instrucción en lenguaje natural.",
    "El asistente responde al momento con texto claro y directo.",
    "Si le pides una acción de agenda, CRM o correo, la realiza y te avisa cuando queda lista.",
]:
    bullet(b)

body("Ejemplos de mensajes:")
for b in [
    "¿Qué clientes tenemos pendientes por llamar hoy?",
    "Agenda una cita para el señor Pérez mañana a las 4 pm.",
    "Envíanos un reporte de los leads de la semana.",
    "¿Qué promociones están activas esta semana?",
    "Busca el correo de Coljuegos del viernes.",
]:
    bullet(b)

body("Notas de voz: puedes enviarlas naturalmente. El asistente las escucha, las entiende y responde por escrito (o de viva voz, si lo pides). Funciona con acento colombiano y ruido de fondo cotidiano.")

heading("4. Límites que debes conocer")
body("Importante: lee esta sección con atención.")
for b in [
    "No inventa datos: si no conoce un dato, te lo dice y te pide que lo confirmes.",
    "No promete premios, pagos ni beneficios fuera de lo que autoriza la regulación de juegos de azar (Coljuegos).",
    "No comparte información del casino con personas no autorizadas: la confidencialidad es lo primero.",
    "Si le adjuntas información sensible (datos de clientes, operaciones de pago), pedirá confirmación antes de realizar la acción.",
    "Atiende los 7 días de la semana, pero las acciones externas (facturar, contactar proveedores) se hacen según el horario del negocio (L–V 9am a 6pm).",
]:
    bullet(b)

heading("5. Reportes y avisos automáticos")
body("El asistente está programado para avisarte y enviarte información sin que tengas que pedirla; solo tienes que configurar el reporte con la información que necesitas, la hora, los días y la frecuencia.")

heading("6. Cómo se gestionan los clientes (CRM)")
body("Proceso interno de NeuralCrew — El asistente registra y organiza los progresos en un embudo de ventas, paso a paso.")
for b in [
    "Lead nuevo (contacto inicial) que llega al sitio web gracias a los anuncios",
    "Juega a la ruleta",
    "Gana el bono",
    "Recibe el email de cómo llegar al casino",
    "Reclama el bono",
    "Nuevo cliente",
]:
    bullet(b)
body("Los clientes que se registren en las páginas web irán a una tabla en Drive para hacer la base de datos.")

heading("7. Los clientes también pueden escribirle")
body("El asistente puede recibir mensajes tanto del equipo del casino como de clientes y distingue el contexto. Su canal de WhatsApp permite:")
for b in [
    "Atención al cliente: resolver dudas de horarios, promociones, ubicaciones y reservas.",
    "Espera en fila: clientes que los escogen fuera del horario de atención, el asistente les contesta siempre y deja un registro para que el equipo retome.",
]:
    bullet(b)
body("Los mensajes de clientes que requieran una decisión humana (pagos, VIP, cobranzas) se escalan al equipo directivo.")
body("Los números y la atención al cliente estarán en las páginas web y publicidad del sector; el asistente se encargará de la atención o, si es necesaria, de la escalación.")

heading("8. Seguridad y privacidad")
for b in [
    "Los datos del casino y de sus clientes son confidenciales y están protegidos.",
    "Solo las personas autorizadas por el Gerente pueden acceder a información sensible.",
    "El asistente opera bajo la Ley 1581 (protección de datos en Colombia) y las normas de Coljuegos sobre juegos de azar.",
]:
    bullet(b)

heading("9. Soporte")
body("NeuralCrew Labs está detrás del asistente. Ante cualquier duda, falla o mejora:")
for b in [
    "WhatsApp principal: 316 691 0728 (Jonathan Parra / administración)",
    "Correo: captain@neuralcrewlabs.com",
]:
    bullet(b)
body("El servicio incluye atención técnica continua, ajustes al asistente y nuevas funcionalidades según las necesidades del sector.")

# Final note
pfin = doc.add_paragraph()
pfin.paragraph_format.space_before = Pt(14); pfin.paragraph_format.space_after = Pt(0)
pfin.alignment = WD_ALIGN_PARAGRAPH.CENTER
gold_bottom_border(pfin, 4)
r = pfin.add_run("Este manual es de uso interno del equipo directivo autorizado.")
r.italic=True; r.font.size=Pt(9); r.font.color.rgb=MGRAY; r.font.name=FONT

# Footer
f = sec.footer
fp = f.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_before = Pt(6)
r = fp.add_run("NeuralCrew Labs   ·   Bogotá   |   3166910728   ·   captain@neuralcrewlabs.com")
r.font.size=Pt(8); r.font.color.rgb=MGRAY; r.font.name=FONT

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("SAVED:", OUT, os.path.getsize(OUT))